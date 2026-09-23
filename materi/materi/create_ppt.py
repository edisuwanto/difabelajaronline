import os
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor

def generate_pptx(html_file, pptx_file):
    with open(html_file, 'r', encoding='utf-8') as f:
        html_content = f.read()

    soup = BeautifulSoup(html_content, 'html.parser')
    main_content = soup.find('main', id='main-content')
    if not main_content:
        main_content = soup.find('body')

    prs = Presentation()
    title_content_layout = prs.slide_layouts[1]
    title_slide_layout = prs.slide_layouts[0]
    
    title = soup.find('title')
    h1 = soup.find('h1')
    main_title = h1.get_text(separator=' ', strip=True) if h1 else (title.get_text() if title else "Materi Machine Learning")
    
    slide = prs.slides.add_slide(title_slide_layout)
    title_shape = slide.shapes.title
    title_shape.text = main_title
    if slide.placeholders[1]:
        slide.placeholders[1].text = "Oleh: Difabelajar\nDeep Learning dalam AI Engineering"
        
    current_slide = None
    tf = None
    current_title = main_title
    lines_on_slide = 0
    MAX_LINES = 13
    
    def add_new_slide(title_text):
        nonlocal current_slide, tf, lines_on_slide, current_title
        current_title = title_text
        current_slide = prs.slides.add_slide(title_content_layout)
        title_shape = current_slide.shapes.title
        title_shape.text = title_text
        body_shape = current_slide.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        tf.word_wrap = True
        lines_on_slide = 0

    def add_paragraph(text, level=0, is_code=False):
        nonlocal current_slide, tf, lines_on_slide, current_title
        if not text.strip():
            return
            
        estimated_lines = text.count('\n') + max(1, len(text) // 60)
        
        if current_slide is None or (lines_on_slide + estimated_lines > MAX_LINES and lines_on_slide > 0):
            add_new_slide(current_title + " (Lanjutan)" if current_slide else current_title)
            
        p = tf.add_paragraph()
        p.text = text
        p.level = level
        if is_code:
            p.font.name = 'Courier New'
            p.font.size = Pt(14)
            p.font.color.rgb = RGBColor(0, 51, 102)
        else:
            p.font.size = Pt(18)
        
        lines_on_slide += estimated_lines

    add_new_slide(main_title)

    for element in main_content.find_all(['h2', 'h3', 'p', 'ul', 'ol', 'pre']):
        if element.name == 'p' and element.find_parent(['ul', 'ol']):
            continue
        if element.name in ['ul', 'ol'] and element.find_parent(['ul', 'ol']):
            continue

        if element.name in ['h2', 'h3']:
            title_text = element.get_text(separator=' ', strip=True)
            if title_text.lower() == 'tujuan pembelajaran' or 'materi' in title_text.lower():
                pass # Still add slide
            add_new_slide(title_text)
            
        elif element.name == 'p':
            text = element.get_text(separator=' ', strip=True)
            if text:
                add_paragraph(text)
                
        elif element.name in ['ul', 'ol']:
            for li in element.find_all('li', recursive=False):
                text = li.get_text(separator=' ', strip=True)
                if text:
                    add_paragraph(text, level=1)
                
        elif element.name == 'pre':
            code_text = element.get_text()
            if code_text:
                add_paragraph(code_text, is_code=True)
                
    prs.save(pptx_file)
    print(f"Berhasil membuat file: {pptx_file}")

if __name__ == "__main__":
    generate_pptx('modulaiadvance4.html', 'modulaiadvance4.pptx')
